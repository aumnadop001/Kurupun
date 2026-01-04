from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from .models import DocumentRecord, Inventory, InventoryTransaction
from urllib.parse import quote

# DocumentRecord data
# {
#     "id": 8,
#     "document_record_inventory_number": "4620-001-0001",
#     "document_record_register_no": "0005-69",
#     "unit_item": "อัน",
#     "first_item": "เครื่องกลั่นน้ำแบบต่าง ๆ",
#     "transactions": [
#         {
#             "id": 22,
#             "transaction_type": "RECEIVE",
#             "transaction_date": "2026-01-04",
#             "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#             "unit_price": "30.00",
#             "type": "INITIAL",
#             "quantity": 300,
#             "total_borrowed": 0,
#             "signature": "",
#             "created_at": "2026-01-04T11:34:52.448454Z",
#             "inventory": 8
#         },
#         {
#             "id": 23,
#             "transaction_type": "RECEIVE",
#             "transaction_date": "2026-01-04",
#             "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#             "unit_price": "30.00",
#             "type": "INITIAL",
#             "quantity": 300,
#             "total_borrowed": 0,
#             "signature": "",
#             "created_at": "2026-01-04T11:34:52.473435Z",
#             "inventory": 8
#         },
#         {
#             "id": 24,
#             "transaction_type": "RECEIVE",
#             "transaction_date": "2026-01-04",
#             "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#             "unit_price": "30.00",
#             "type": "INITIAL",
#             "quantity": 300,
#             "total_borrowed": 0,
#             "signature": "",
#             "created_at": "2026-01-04T11:34:52.495693Z",
#             "inventory": 8
#         },
#         {
#             "id": 25,
#             "transaction_type": "ISSUE",
#             "transaction_date": "2026-01-04",
#             "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#             "unit_price": "30.00",
#             "type": "INITIAL",
#             "quantity": 300,
#             "total_borrowed": 200,
#             "signature": "",
#             "created_at": "2026-01-04T11:34:52.524627Z",
#             "inventory": 8
#         },
#         {
#             "id": 26,
#             "transaction_type": "ISSUE",
#             "transaction_date": "2026-01-04",
#             "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#             "unit_price": "30.00",
#             "type": "INITIAL",
#             "quantity": 100,
#             "total_borrowed": 0,
#             "signature": "",
#             "created_at": "2026-01-04T11:34:52.550323Z",
#             "inventory": 8
#         }
#     ],
#     "pending_date": null,
#     "pending_evidence": "",
#     "pending_unit": "",
#     "pending_quantity": null,
#     "pending_receive1": null,
#     "pending_balance1": null,
#     "pending_receive2": null,
#     "pending_balance2": null,
#     "pending_receive3": null,
#     "pending_balance3": null,
#     "pending_receive4": null,
#     "pending_balance4": null,
#     "pending_signature": "",
#     "request_date": null,
#     "received_quantity": 0,
#     "unit_price": "0.00",
#     "request_evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#     "request_type": "INITIAL",
#     "issue_quantity": 0,
#     "total_borrowed": 0,
#     "previous_stock_balance": 0,
#     "stock_balance": 300,
#     "request_signature": "",
#     "document_record": 5
# }

# Transaction data
# [
#     {
#         "id": 22,
#         "transaction_type": "RECEIVE",
#         "transaction_date": "2026-01-04",
#         "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#         "unit_price": "30.00",
#         "type": "INITIAL",
#         "quantity": 300,
#         "total_borrowed": 0,
#         "signature": "",
#         "created_at": "2026-01-04T11:34:52.448454Z",
#         "inventory": 8
#     },
#     {
#         "id": 23,
#         "transaction_type": "RECEIVE",
#         "transaction_date": "2026-01-04",
#         "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#         "unit_price": "30.00",
#         "type": "INITIAL",
#         "quantity": 300,
#         "total_borrowed": 0,
#         "signature": "",
#         "created_at": "2026-01-04T11:34:52.473435Z",
#         "inventory": 8
#     },
#     {
#         "id": 24,
#         "transaction_type": "RECEIVE",
#         "transaction_date": "2026-01-04",
#         "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#         "unit_price": "30.00",
#         "type": "INITIAL",
#         "quantity": 300,
#         "total_borrowed": 0,
#         "signature": "",
#         "created_at": "2026-01-04T11:34:52.495693Z",
#         "inventory": 8
#     },
#     {
#         "id": 25,
#         "transaction_type": "ISSUE",
#         "transaction_date": "2026-01-04",
#         "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#         "unit_price": "30.00",
#         "type": "INITIAL",
#         "quantity": 300,
#         "total_borrowed": 200,
#         "signature": "",
#         "created_at": "2026-01-04T11:34:52.524627Z",
#         "inventory": 8
#     },
#     {
#         "id": 26,
#         "transaction_type": "ISSUE",
#         "transaction_date": "2026-01-04",
#         "evidence": "ใบแจ้งหนี้ - งานบริหารงานทั่วไป - 0005-69",
#         "unit_price": "30.00",
#         "type": "INITIAL",
#         "quantity": 100,
#         "total_borrowed": 0,
#         "signature": "",
#         "created_at": "2026-01-04T11:34:52.550323Z",
#         "inventory": 8
#     }
# ]
def export_document_record_to_excel(document_record):
    """
    Export DocumentRecord และ Inventories ไปเป็นไฟล์ Excel
    ตามโครงสร้างของ HTML template
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"ทะเบียนคุม_{document_record.registration_number}"

    # Define styles
    header_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    gray_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    # Set column widths (ปรับให้พอดีกับ A4)
    column_widths = {
        "A": 12,   # ว.ด.ป.
        "B": 14,   # หลักฐาน
        "C": 10,   # หน่วย
        "D": 12,   # จำนวน
        "E": 14,   # รับ/ค้าง 1
        "F": 14,   # รับ/ค้าง 2
        "G": 14,   # รับ/ค้าง 3
        "H": 14,   # รับ/ค้าง 4
        "I": 12,   # ว.ด.ป.
        "J": 10,   # รับ
        "K": 10,   # ราคาต่อหน่วย
        "L": 18,   # หลักฐาน
        "M": 10,   # ขั้นต้น
        "N": 10,   # ทดแทน
        "O": 10,   # จ่าย
        "P": 10,   # รวมยืม
        "Q": 12,   # คงคลัง
        "R": 14,   # ลายมือชื่อ
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Row 1-2: Header section
    # Row 1
    ws.merge_cells("A1:A2")
    ws["A1"] = "ว.ด.ป."
    ws["A1"].font = header_font
    ws["A1"].fill = gray_fill
    ws["A1"].alignment = center_alignment
    ws["A1"].border = thin_border

    ws.merge_cells("B1:B2")
    ws["B1"] = (
        document_record.registration_date.strftime("%d/%m/%Y")
        if document_record.registration_date
        else ""
    )
    ws["B1"].font = header_font
    ws["B1"].fill = gray_fill
    ws["B1"].alignment = center_alignment
    ws["B1"].border = thin_border

    ws.merge_cells("C1:D1")
    ws["C1"] = "หมายเลขพัสดุ"
    ws["C1"].font = header_font
    ws["C1"].fill = gray_fill
    ws["C1"].alignment = center_alignment
    ws["C1"].border = thin_border

    ws.merge_cells("E1:H1")
    ws["E1"] = document_record.inventory_number or ""
    ws["E1"].font = header_font
    ws["E1"].fill = gray_fill
    ws["E1"].alignment = center_alignment
    ws["E1"].border = thin_border

    ws.merge_cells("I1:N2")
    ws["I1"] = "รายการ"
    ws["I1"].font = header_font
    ws["I1"].fill = gray_fill
    ws["I1"].alignment = center_alignment
    ws["I1"].border = thin_border

    ws.merge_cells("O1:O2")
    ws["O1"] = "หน่วยนับ"
    ws["O1"].font = header_font
    ws["O1"].fill = gray_fill
    ws["O1"].alignment = center_alignment
    ws["O1"].border = thin_border

    ws.merge_cells("P1:P2")
    ws["P1"] = document_record.unit_of_measure or ""
    ws["P1"].font = header_font
    ws["P1"].fill = gray_fill
    ws["P1"].alignment = center_alignment
    ws["P1"].border = thin_border

    ws.merge_cells("Q1:R2")
    ws["Q1"] = f"ที่เก็บ {document_record.storage_location or ''}"
    ws["Q1"].font = header_font
    ws["Q1"].fill = gray_fill
    ws["Q1"].alignment = center_alignment
    ws["Q1"].border = thin_border

    # Row 2
    ws["C2"] = "วัน"
    ws["C2"].font = header_font
    ws["C2"].fill = gray_fill
    ws["C2"].alignment = center_alignment
    ws["C2"].border = thin_border

    ws["D2"] = "จำนวน"
    ws["D2"].font = header_font
    ws["D2"].fill = gray_fill
    ws["D2"].alignment = center_alignment
    ws["D2"].border = thin_border

    ws.merge_cells("E2:H2")
    ws["E2"] = "หมายเลขพัสดุแทนกันได้"
    ws["E2"].font = header_font
    ws["E2"].fill = gray_fill
    ws["E2"].alignment = center_alignment
    ws["E2"].border = thin_border

    # Row 3-5: Criteria section
    ws.merge_cells("A3:B3")
    ws["A3"] = "เกณฑ์สั่ง"
    ws["A3"].font = header_font
    ws["A3"].fill = gray_fill
    ws["A3"].alignment = center_alignment
    ws["A3"].border = thin_border

    ws["C3"] = document_record.days_to_order or ""
    ws["C3"].border = thin_border
    ws["C3"].alignment = center_alignment

    ws["D3"] = document_record.quantity_to_order or ""
    ws["D3"].border = thin_border
    ws["D3"].alignment = center_alignment

    ws.merge_cells("E3:H5")
    alternate_numbers = document_record.inventory_alternate_numbers
    if isinstance(alternate_numbers, list):
        ws["E3"] = ", ".join(alternate_numbers)
    elif alternate_numbers:
        ws["E3"] = str(alternate_numbers)
    ws["E3"].border = thin_border
    ws["E3"].alignment = left_alignment

    ws.merge_cells("I3:N5")
    equipment = document_record.related_equipment
    equipment_text = "ครุภัณฑ์ที่เกี่ยวข้อง: "
    if isinstance(equipment, list):
        equipment_text += ", ".join(equipment)
    elif equipment:
        equipment_text += str(equipment)
    ws["I3"] = equipment_text
    ws["I3"].font = header_font
    ws["I3"].fill = gray_fill
    ws["I3"].border = thin_border
    ws["I3"].alignment = left_alignment

    ws.merge_cells("O3:R5")
    ws["O3"] = f"หมายเหตุ: {document_record.remark or ''}"
    ws["O3"].font = header_font
    ws["O3"].fill = gray_fill
    ws["O3"].border = thin_border
    ws["O3"].alignment = left_alignment

    # Row 4
    ws.merge_cells("A4:B4")
    ws["A4"] = "จุดสั่งเพิ่มเติม"
    ws["A4"].font = header_font
    ws["A4"].fill = gray_fill
    ws["A4"].alignment = center_alignment
    ws["A4"].border = thin_border

    ws["C4"] = document_record.reorder_point_days or ""
    ws["C4"].border = thin_border
    ws["C4"].alignment = center_alignment

    ws["D4"] = document_record.reorder_point_quantity or ""
    ws["D4"].border = thin_border
    ws["D4"].alignment = center_alignment

    # Row 5
    ws.merge_cells("A5:B5")
    ws["A5"] = "เกณฑ์ปลอดภัย"
    ws["A5"].font = header_font
    ws["A5"].fill = gray_fill
    ws["A5"].alignment = center_alignment
    ws["A5"].border = thin_border

    ws["C5"] = document_record.safety_stock_days or ""
    ws["C5"].border = thin_border
    ws["C5"].alignment = center_alignment

    ws["D5"] = document_record.safety_stock_quantity or ""
    ws["D5"].border = thin_border
    ws["D5"].alignment = center_alignment

    # Row 6: Section headers
    ws.merge_cells("A6:H6")
    ws["A6"] = "ค้างรับ และ ค้างจ่าย"
    ws["A6"].font = header_font
    ws["A6"].fill = gray_fill
    ws["A6"].alignment = center_alignment
    ws["A6"].border = thin_border

    ws.merge_cells("I6:R6")
    ws["I6"] = "ความต้องการรับและจ่าย"
    ws["I6"].font = header_font
    ws["I6"].fill = gray_fill
    ws["I6"].alignment = center_alignment
    ws["I6"].border = thin_border

    # Row 7-8: Column headers
    headers_row7 = [
        ("A7:A8", "ว.ด.ป."),
        ("B7:B8", "หลักฐาน"),
        ("C7:C8", "หน่วย"),
        ("D7:D8", "จำนวน"),
        ("E7:E8", "รับ\nค้าง"),
        ("F7:F8", "รับ\nค้าง"),
        ("G7:G8", "รับ\nค้าง"),
        ("H7:H8", "รับ\nค้าง"),
        ("I7:I8", "ว.ด.ป."),
        ("J7:J8", "รับ"),
        ("K7:K8", "ราคาต่อหน่วย"),
        ("L7:L8", "หลักฐาน"),
        ("O7:O8", "จ่าย"),
        ("P7:P8", "รวมยืม"),
        ("Q7:Q8", "คงคลัง"),
        ("R7:R8", "ลายมือชื่อ"),
    ]

    for cell_range, text in headers_row7:
        if ":" in cell_range:
            ws.merge_cells(cell_range)
            cell = cell_range.split(":")[0]
        else:
            cell = cell_range
        ws[cell] = text
        ws[cell].font = header_font
        ws[cell].fill = gray_fill
        ws[cell].alignment = center_alignment
        ws[cell].border = thin_border

    ws.merge_cells("M7:N7")
    ws["M7"] = "ความต้องการ"
    ws["M7"].font = header_font
    ws["M7"].fill = gray_fill
    ws["M7"].alignment = center_alignment
    ws["M7"].border = thin_border

    # Row 8: Sub-headers
    ws["M8"] = "ขั้นต้น"
    ws["M8"].font = header_font
    ws["M8"].fill = gray_fill
    ws["M8"].alignment = center_alignment
    ws["M8"].border = thin_border

    ws["N8"] = "ทดแทน"
    ws["N8"].font = header_font
    ws["N8"].fill = gray_fill
    ws["N8"].alignment = center_alignment
    ws["N8"].border = thin_border

    # Row 9+: Data rows (InventoryTransactions)
    # ดึง transactions จาก inventory ที่เกี่ยวข้องกับ document_record
    current_row = 9
    cumulative_received = 0  # ยอดรับสะสม
    cumulative_issued = 0    # ยอดจ่ายสะสม
    cumulative_borrowed = 0  # ยอดยืมสะสม
    
    # ดึง inventory ของ document_record นี้
    try:
        inventory = document_record.inventory
        # ดึง transactions ทั้งหมดที่เกี่ยวข้อง เรียงตามวันที่
        transactions = inventory.transactions.all().order_by('transaction_date', 'created_at')
        
        # เริ่มต้นด้วย previous_stock_balance ถ้ามี
        previous_balance = inventory.previous_stock_balance or 0
        
        for transaction in transactions:
            # Pending section - ใช้ข้อมูลจาก inventory header (ถ้ามี)
            if transaction == transactions.first():
                # แสดง pending info ในแถวแรกเท่านั้น
                ws[f"A{current_row}"] = (
                    inventory.pending_date.strftime("%d/%m/%Y")
                    if inventory.pending_date
                    else ""
                )
                ws[f"B{current_row}"] = inventory.pending_evidence or ""
                ws[f"C{current_row}"] = inventory.pending_unit or ""
                ws[f"D{current_row}"] = inventory.pending_quantity or ""

                # Pending receive/balance columns (4 sets)
                pending_1 = (
                    f"{inventory.pending_receive1 or ''}/{inventory.pending_balance1 or ''}"
                )
                ws[f"E{current_row}"] = (
                    pending_1
                    if inventory.pending_receive1 or inventory.pending_balance1
                    else ""
                )

                pending_2 = (
                    f"{inventory.pending_receive2 or ''}/{inventory.pending_balance2 or ''}"
                )
                ws[f"F{current_row}"] = (
                    pending_2
                    if inventory.pending_receive2 or inventory.pending_balance2
                    else ""
                )

                pending_3 = (
                    f"{inventory.pending_receive3 or ''}/{inventory.pending_balance3 or ''}"
                )
                ws[f"G{current_row}"] = (
                    pending_3
                    if inventory.pending_receive3 or inventory.pending_balance3
                    else ""
                )

                pending_4 = (
                    f"{inventory.pending_receive4 or ''}/{inventory.pending_balance4 or ''}"
                )
                ws[f"H{current_row}"] = (
                    pending_4
                    if inventory.pending_receive4 or inventory.pending_balance4
                    else ""
                )
            else:
                # แถวอื่นๆ ไม่แสดง pending info
                ws[f"A{current_row}"] = ""
                ws[f"B{current_row}"] = ""
                ws[f"C{current_row}"] = ""
                ws[f"D{current_row}"] = ""
                ws[f"E{current_row}"] = ""
                ws[f"F{current_row}"] = ""
                ws[f"G{current_row}"] = ""
                ws[f"H{current_row}"] = ""

            # Request section - ข้อมูลจาก transaction
            ws[f"I{current_row}"] = (
                transaction.transaction_date.strftime("%d/%m/%Y")
                if transaction.transaction_date
                else ""
            )
            
            # คำนวณยอดสะสม
            if transaction.transaction_type == 'RECEIVE':
                cumulative_received += transaction.quantity or 0
                ws[f"J{current_row}"] = transaction.quantity or 0
                ws[f"O{current_row}"] = ""  # ไม่มีการจ่าย
                ws[f"P{current_row}"] = ""  # ไม่มีการยืม
            else:  # ISSUE
                cumulative_issued += transaction.quantity or 0
                cumulative_borrowed += transaction.total_borrowed or 0
                ws[f"J{current_row}"] = ""  # ไม่มีการรับ
                ws[f"O{current_row}"] = transaction.quantity or 0
                ws[f"P{current_row}"] = transaction.total_borrowed or 0
            
            ws[f"K{current_row}"] = (
                float(transaction.unit_price) if transaction.unit_price else ""
            )
            ws[f"L{current_row}"] = transaction.evidence or ""

            # Request type checkboxes
            ws[f"M{current_row}"] = "✓" if transaction.type == "INITIAL" else ""
            ws[f"N{current_row}"] = "✓" if transaction.type == "REPLACEMENT" else ""

            # คงคลัง = ยอดเริ่มต้น + รับสะสม - จ่ายสะสม - ยืมสะสม
            calculated_stock_balance = previous_balance + cumulative_received - cumulative_issued - cumulative_borrowed
            ws[f"Q{current_row}"] = calculated_stock_balance
            ws[f"R{current_row}"] = transaction.signature or ""

            # Apply styles to all cells in the row
            for col in range(1, 19):  # A to R
                cell = ws.cell(row=current_row, column=col)
                cell.font = normal_font
                cell.alignment = center_alignment
                cell.border = thin_border

            current_row += 1
            
    except Inventory.DoesNotExist:
        # ถ้าไม่มี inventory ก็ข้ามไป
        pass

    # Set row heights (ปรับให้พอดีกับ A4)
    for row in range(1, current_row):
        ws.row_dimensions[row].height = 16
    
    # Set page setup for A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    
    # Set margins
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    return wb


def export_document_records_table_to_excel(document_records):
    """
    Export DocumentRecords เป็น Excel ตามโครงสร้างของ table.html
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "ทะเบียนเอกสาร"

    # Define styles
    header_font = Font(name="TH SarabunPSK", size=14, bold=True)
    normal_font = Font(name="TH SarabunPSK", size=14)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    
    header_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

    # Set column widths
    column_widths = {
        "A": 15,  # ทะเบียนที่
        "B": 18,  # วันที่ลงทะเบียน
        "C": 15,  # เอกสาร
        "D": 20,  # จาก
        "E": 25,  # รายการแรกในเอกสาร
        "F": 20,  # (ทะเบียนที่ - duplicate in HTML)
        "G": 18,  # วันที่เก็บเข้าแฟ้ม
        "H": 25,  # เลขที่เอกสารที่เกี่ยวข้อง
        "I": 20,  # ชื่อผู้เบิก
        "J": 18,  # เลขที่ชุดเบิก
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Header row
    headers = [
        "ทะเบียนที่",
        "วันที่ลงทะเบียน",
        "เอกสาร",
        "จาก",
        "รายการแรกในเอกสาร",
        "",  # colspan 2 in HTML
        "วันที่เก็บเข้าแฟ้ม",
        "เลขที่เอกสารที่เกี่ยวข้อง",
        "ชื่อผู้เบิก",
        "เลขที่ชุดเบิก",
    ]
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
        cell.fill = header_fill

    # Data rows
    current_row = 2
    for doc in document_records:
        ws.cell(row=current_row, column=1).value = doc.registerNo or ""
        ws.cell(row=current_row, column=2).value = (
            doc.registration_date.strftime("%d/%m/%Y") if doc.registration_date else ""
        )
        ws.cell(row=current_row, column=3).value = doc.document_type or ""
        ws.cell(row=current_row, column=4).value = doc.sender or ""
        ws.cell(row=current_row, column=5).value = doc.first_item or ""
        ws.cell(row=current_row, column=6).value = doc.registration_number or ""
        ws.cell(row=current_row, column=7).value = (
            doc.file_storage_date.strftime("%d/%m/%Y") if doc.file_storage_date else ""
        )
        ws.cell(row=current_row, column=8).value = doc.related_document_number or ""
        ws.cell(row=current_row, column=9).value = doc.requester_name or ""
        ws.cell(row=current_row, column=10).value = doc.requester_set_number or ""

        # Apply styles to all cells in the row
        for col in range(1, 11):
            cell = ws.cell(row=current_row, column=col)
            cell.font = normal_font
            cell.alignment = center_alignment
            cell.border = thin_border

        current_row += 1

    # Set row heights
    ws.row_dimensions[1].height = 25
    for row in range(2, current_row):
        ws.row_dimensions[row].height = 20

    return wb


def generate_table_excel_response(document_records):
    """
    สร้าง Excel response สำหรับตาราง DocumentRecords
    """
    wb = export_document_records_table_to_excel(document_records)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    filename = "ทะเบียนเอกสาร.xlsx"
    quoted_filename = quote(filename)

    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quoted_filename}"

    wb.save(response)
    return response


def generate_excel_response(document_record):
    wb = export_document_record_to_excel(document_record)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    filename = f"ทะเบียนคุมวัสดุ_{document_record.registration_number}.xlsx"
    quoted_filename = quote(filename)

    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quoted_filename}"

    wb.save(response)
    return response
