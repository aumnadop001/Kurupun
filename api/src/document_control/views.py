from django.shortcuts import render
from django.http import HttpResponse
from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django_filters.rest_framework import DjangoFilterBackend
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
from decimal import Decimal
from rest_framework.pagination import PageNumberPagination
from src.document_control.models import (
    Item,
    AlternativeItem,
    RelatedEquipment,
    PendingTransaction,
    StockTransaction,
    Description
)
from src.document_control.serializers import (
    ItemSerializer,
    ItemListSerializer,
    AlternativeItemSerializer,
    RelatedEquipmentSerializer,
    PendingTransactionSerializer,
    StockTransactionSerializer,
    StockCardExportSerializer,
    DescriptionNestedSerializer,
    DescriptionSerializer
)

class StandardResultsSetPagination(PageNumberPagination):
    page_size = 100
    page_size_query_param = 'page_size'
    max_page_size = 1000

class ItemViewSet(viewsets.ModelViewSet):
    """
    ViewSet for Item CRUD operations
    """

    queryset = Item.objects.select_related(
        "description",
        "description__class_id",
        "description__type_id",
        "description__gpsc_id",
    ).all()
    serializer_class = ItemSerializer
    # permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ["is_active", "storage_location"]
    search_fields = [
        "description__Des_name",
        "description__Des_id",
        "description__class_id__class_name",
        "description__type_id__ptype_name",
        "storage_location",
    ]
    ordering_fields = ["created_at", "updated_at"]
    ordering = ["-created_at"]
    pagination_class = StandardResultsSetPagination
    def get_serializer_class(self):
        if self.action == "list":
            return ItemListSerializer
        return ItemSerializer

    def create(self, request, *args, **kwargs):
        """Create item with alternatives and related equipments"""
        alternatives_data = request.data.pop('alternatives', [])
        equipments_data = request.data.pop('equipments', [])
        pending_transactions_data = request.data.pop('pending_transactions', [])
        stock_transactions_data = request.data.pop('stock_transactions', [])
        
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        item = serializer.save()
        
        # Create alternatives
        for alt_item_id in alternatives_data:
            AlternativeItem.objects.create(item=item, alternative_item_id=alt_item_id)
        
        # Create related equipments
        for equipment in equipments_data:
            RelatedEquipment.objects.create(
                item=item,
                equipment_name=equipment.get('equipment_name', ''),
                equipment_code=equipment.get('equipment_code', '')
            )
        
        # Create pending transactions
        for pt in pending_transactions_data:
            PendingTransaction.objects.create(
                item=item,
                transaction_type=pt.get('transaction_type', 'receive'),
                date=pt.get('date'),
                document_no=pt.get('document_no', ''),
                unit=pt.get('unit', ''),
                quantity=pt.get('quantity', 0)
            )
        
        # Create stock transactions
        for st in stock_transactions_data:
            StockTransaction.objects.create(
                item=item,
                date=st.get('date'),
                receive_quantity=st.get('receive_quantity', 0),
                unit_price=st.get('unit_price', 0),
                receive_document_no=st.get('receive_document_no', ''),
                initial_demand=st.get('initial_demand', 0),
                replacement_demand=st.get('replacement_demand', 0),
                issue_quantity=st.get('issue_quantity', 0),
                total_borrowed=st.get('total_borrowed', 0),
                signature=st.get('signature', ''),
                notes=st.get('notes', '')
            )
        
        headers = self.get_success_headers(serializer.data)
        return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)
    
    def update(self, request, *args, **kwargs):
        """Update item with alternatives and related equipments"""
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        
        alternatives_data = request.data.pop('alternatives', None)
        equipments_data = request.data.pop('equipments', None)
        pending_transactions_data = request.data.pop('pending_transactions', None)
        stock_transactions_data = request.data.pop('stock_transactions', None)
        
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        
        # Update alternatives if provided
        if alternatives_data is not None:
            # Delete existing alternatives
            AlternativeItem.objects.filter(item=instance).delete()
            # Create new ones
            for alt_item_id in alternatives_data:
                AlternativeItem.objects.create(item=instance, alternative_item_id=alt_item_id)
        
        # Update equipments if provided
        if equipments_data is not None:
            # Delete existing equipments
            RelatedEquipment.objects.filter(item=instance).delete()
            # Create new ones
            for equipment in equipments_data:
                RelatedEquipment.objects.create(
                    item=instance,
                    equipment_name=equipment.get('equipment_name', ''),
                    equipment_code=equipment.get('equipment_code', '')
                )
        
        # Update pending transactions if provided
        if pending_transactions_data is not None:
            # Delete existing pending transactions
            PendingTransaction.objects.filter(item=instance).delete()
            # Create new ones
            for pt in pending_transactions_data:
                PendingTransaction.objects.create(
                    item=instance,
                    transaction_type=pt.get('transaction_type', 'receive'),
                    date=pt.get('date'),
                    document_no=pt.get('document_no', ''),
                    unit=pt.get('unit', ''),
                    quantity=pt.get('quantity', 0)
                )
        
        # Update stock transactions if provided
        if stock_transactions_data is not None:
            # Delete existing stock transactions
            StockTransaction.objects.filter(item=instance).delete()
            # Create new ones
            for st in stock_transactions_data:
                StockTransaction.objects.create(
                    item=instance,
                    date=st.get('date'),
                    receive_quantity=st.get('receive_quantity', 0),
                    unit_price=st.get('unit_price', 0),
                    receive_document_no=st.get('receive_document_no', ''),
                    initial_demand=st.get('initial_demand', 0),
                    replacement_demand=st.get('replacement_demand', 0),
                    issue_quantity=st.get('issue_quantity', 0),
                    total_borrowed=st.get('total_borrowed', 0),
                    signature=st.get('signature', ''),
                    notes=st.get('notes', '')
                )
        
        if getattr(instance, '_prefetched_objects_cache', None):
            instance._prefetched_objects_cache = {}
        
        return Response(serializer.data)

    @action(detail=True, methods=["get"], url_path="stock-card")
    def stock_card(self, request, pk=None):
        """Get stock card data for specific item"""
        item = self.get_object()
        transactions = StockTransaction.objects.filter(item=item).order_by("date")
        alternatives = AlternativeItem.objects.filter(item=item)
        equipments = RelatedEquipment.objects.filter(item=item)
        pending = PendingTransaction.objects.filter(item=item, is_resolved=False)

        return Response(
            {
                "item": ItemSerializer(item).data,
                "transactions": StockTransactionSerializer(transactions, many=True).data,
                "alternatives": AlternativeItemSerializer(alternatives, many=True).data,
                "equipments": RelatedEquipmentSerializer(equipments, many=True).data,
                "pending": PendingTransactionSerializer(pending, many=True).data,
            }
        )

    @action(detail=True, methods=["get"], url_path="export-excel")
    def export_excel(self, request, pk=None):
        """Export stock card to Excel file"""
        item = self.get_object()
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")

        # Filter transactions by date if provided
        transactions = StockTransaction.objects.filter(item=item)
        if start_date:
            transactions = transactions.filter(date__gte=start_date)
        if end_date:
            transactions = transactions.filter(date__lte=end_date)
        transactions = transactions.order_by("date")

        # Get related data
        alternatives = AlternativeItem.objects.filter(item=item)
        equipments = RelatedEquipment.objects.filter(item=item)
        pending_transactions = PendingTransaction.objects.filter(item=item).order_by("date")

        # Create Excel file
        wb = Workbook()
        ws = wb.active
        ws.title = "บัตรคุมพัสดุ"

        # Styles
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Header section - Row 1-2
        ws.merge_cells("A1:B2")
        ws["A1"] = "ว.ด.ป."
        ws["A1"].fill = header_fill
        ws["A1"].font = bold_font
        ws["A1"].alignment = center_align
        ws["A1"].border = border

        ws.merge_cells("C1:D1")
        ws["C1"] = "หมายเลขพัสดุ"
        ws["C1"].fill = header_fill
        ws["C1"].font = bold_font
        ws["C1"].alignment = center_align
        ws["C1"].border = border

        ws.merge_cells("E1:H1")
        ws["E1"] = item.get_item_id
        ws["E1"].alignment = center_align
        ws["E1"].border = border

        ws.merge_cells("I1:N2")
        ws["I1"] = f"รายการ {item.name}"
        ws["I1"].fill = header_fill
        ws["I1"].font = bold_font
        ws["I1"].alignment = center_align
        ws["I1"].border = border

        ws.merge_cells("O1:P2")
        ws["O1"] = f"หน่วยนับ {item.unit}"
        ws["O1"].fill = header_fill
        ws["O1"].font = bold_font
        ws["O1"].alignment = center_align
        ws["O1"].border = border

        ws.merge_cells("Q1:R2")
        ws["Q1"] = f"ที่เก็บ {item.storage_location}"
        ws["Q1"].fill = header_fill
        ws["Q1"].font = bold_font
        ws["Q1"].alignment = center_align
        ws["Q1"].border = border

        # Row 2 - Day and Quantity, Alternatives
        ws["C2"] = "วัน"
        ws["C2"].fill = header_fill
        ws["C2"].font = bold_font
        ws["C2"].alignment = center_align
        ws["C2"].border = border

        ws["D2"] = "จำนวน"
        ws["D2"].fill = header_fill
        ws["D2"].font = bold_font
        ws["D2"].alignment = center_align
        ws["D2"].border = border

        ws.merge_cells("E2:H2")
        ws["E2"] = "หมายเลขพัสดุแทนกันได้"
        ws["E2"].fill = header_fill
        ws["E2"].font = bold_font
        ws["E2"].alignment = center_align
        ws["E2"].border = border

        # Row 3-5 - Criteria
        ws.merge_cells("A3:B3")
        ws["A3"] = "เกณฑ์สั่ง"
        ws["A3"].fill = header_fill
        ws["A3"].font = bold_font
        ws["A3"].alignment = center_align
        ws["A3"].border = border

        ws["C3"] = item.order_criteria_day
        ws["C3"].alignment = center_align
        ws["C3"].border = border

        ws["D3"] = float(item.order_criteria_quantity)
        ws["D3"].alignment = center_align
        ws["D3"].border = border

        ws.merge_cells("E3:H5")
        alt_text = ", ".join([alt.alternative_item.get_item_id for alt in alternatives])
        ws["E3"] = alt_text
        ws["E3"].alignment = center_align
        ws["E3"].border = border

        ws.merge_cells("I3:N5")
        equip_text = ", ".join([f"{eq.equipment_name}" for eq in equipments])
        ws["I3"] = f"ครุภัณฑ์ที่เกี่ยวข้อง {equip_text}"
        ws["I3"].fill = header_fill
        ws["I3"].font = bold_font
        ws["I3"].alignment = center_align
        ws["I3"].border = border

        ws.merge_cells("O3:R5")
        ws["O3"] = f"หมายเหตุ {item.notes or ''}"
        ws["O3"].fill = header_fill
        ws["O3"].font = bold_font
        ws["O3"].alignment = center_align
        ws["O3"].border = border

        # Row 4
        ws.merge_cells("A4:B4")
        ws["A4"] = "จุดสั่งเพิ่มเติม"
        ws["A4"].fill = header_fill
        ws["A4"].font = bold_font
        ws["A4"].alignment = center_align
        ws["A4"].border = border

        ws["C4"] = item.additional_order_point_day
        ws["C4"].alignment = center_align
        ws["C4"].border = border

        ws["D4"] = float(item.additional_order_point_quantity)
        ws["D4"].alignment = center_align
        ws["D4"].border = border

        # Row 5
        ws.merge_cells("A5:B5")
        ws["A5"] = "เกณฑ์ปลอดภัย"
        ws["A5"].fill = header_fill
        ws["A5"].font = bold_font
        ws["A5"].alignment = center_align
        ws["A5"].border = border

        ws["C5"] = item.safety_criteria_day
        ws["C5"].alignment = center_align
        ws["C5"].border = border

        ws["D5"] = float(item.safety_criteria_quantity)
        ws["D5"].alignment = center_align
        ws["D5"].border = border

        # Row 6 - Section headers
        ws.merge_cells("A6:H6")
        ws["A6"] = "ค้างรับ และ ค้างจ่าย"
        ws["A6"].fill = header_fill
        ws["A6"].font = bold_font
        ws["A6"].alignment = center_align
        ws["A6"].border = border

        ws.merge_cells("I6:R6")
        ws["I6"] = "ความต้องการรับและจ่าย"
        ws["I6"].fill = header_fill
        ws["I6"].font = bold_font
        ws["I6"].alignment = center_align
        ws["I6"].border = border

        # Row 7-8 - Transaction table headers
        headers_row7 = [
            ("A7:A8", "ว.ด.ป."),
            ("B7:B8", "หลักฐาน"),
            ("C7:C8", "หน่วย"),
            ("D7:D8", "จำนวน"),
            ("E7:E8", "รับค้าง"),
            ("F7:F8", "รับค้าง"),
            ("G7:G8", "รับค้าง"),
            ("H7:H8", "รับค้าง"),
            ("I7:I8", "ว.ด.ป."),
            ("J7:J8", "รับ"),
            ("K7:K8", "ราคาต่อหน่วย"),
            ("L7:L8", "หลักฐาน"),
            ("M7:N7", "ความต้องการ"),
            ("O7:O8", "จ่าย"),
            ("P7:P8", "รวมยืม"),
            ("Q7:Q8", "คงคลัง"),
            ("R7:R8", "ลายมือชื่อ"),
        ]

        for cell_range, text in headers_row7:
            ws.merge_cells(cell_range)
            cell = ws[cell_range.split(":")[0]]
            cell.value = text
            cell.fill = header_fill
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = border

        ws["M8"] = "ขั้นต้น"
        ws["M8"].fill = header_fill
        ws["M8"].font = bold_font
        ws["M8"].alignment = center_align
        ws["M8"].border = border

        ws["N8"] = "ทดแทน"
        ws["N8"].fill = header_fill
        ws["N8"].font = bold_font
        ws["N8"].alignment = center_align
        ws["N8"].border = border

        # Data rows - First add pending transactions
        row = 9
        for pt in pending_transactions:
            ws[f"A{row}"] = pt.date.strftime("%d/%m/%Y")
            ws[f"B{row}"] = pt.document_no or ""
            ws[f"C{row}"] = pt.unit or ""
            ws[f"D{row}"] = float(pt.quantity) if pt.quantity else ""
            ws[f"E{row}"] = item.pending_received_1 or ""
            ws[f"F{row}"] = item.pending_received_2 or ""
            ws[f"G{row}"] = item.pending_received_3 or ""
            ws[f"H{row}"] = item.pending_received_4 or ""
            ws[f"I{row}"] = ""
            ws[f"J{row}"] = ""
            ws[f"K{row}"] = ""
            ws[f"L{row}"] = ""
            ws[f"M{row}"] = ""
            ws[f"N{row}"] = ""
            ws[f"O{row}"] = ""
            ws[f"P{row}"] = ""
            ws[f"Q{row}"] = ""
            ws[f"R{row}"] = ""

            # Apply borders and alignment
            for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R"]:
                cell = ws[f"{col}{row}"]
                cell.border = border
                cell.alignment = center_align

            row += 1

        # Then add stock transactions
        for trans in transactions:
            ws[f"A{row}"] = ""
            ws[f"B{row}"] = ""
            ws[f"C{row}"] = ""
            ws[f"D{row}"] = ""
            ws[f"E{row}"] = item.pending_received_1 or ""
            ws[f"F{row}"] = item.pending_received_2 or ""
            ws[f"G{row}"] = item.pending_received_3 or ""
            ws[f"H{row}"] = item.pending_received_4 or ""
            ws[f"I{row}"] = trans.date.strftime("%d/%m/%Y")
            ws[f"J{row}"] = float(trans.receive_quantity) if trans.receive_quantity else ""
            ws[f"K{row}"] = float(trans.unit_price) if trans.unit_price else ""
            ws[f"L{row}"] = trans.receive_document_no or ""
            ws[f"M{row}"] = float(trans.initial_demand) if trans.initial_demand else ""
            ws[f"N{row}"] = float(trans.replacement_demand) if trans.replacement_demand else ""
            ws[f"O{row}"] = float(trans.issue_quantity) if trans.issue_quantity else ""
            ws[f"P{row}"] = float(trans.total_borrowed) if trans.total_borrowed else ""
            ws[f"Q{row}"] = float(trans.stock_balance)
            ws[f"R{row}"] = trans.signature or ""

            # Apply borders and alignment
            for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R"]:
                cell = ws[f"{col}{row}"]
                cell.border = border
                cell.alignment = center_align

            row += 1

        # Adjust column widths
        column_widths = {
            "A": 12, "B": 15, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10, "H": 10,
            "I": 12, "J": 10, "K": 12, "L": 15, "M": 10, "N": 10, "O": 10, "P": 10,
            "Q": 12, "R": 15,
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # Create response
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"stock_card_{item.get_item_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


class AlternativeItemViewSet(viewsets.ModelViewSet):
    """ViewSet for AlternativeItem CRUD operations"""
    pagination_class = StandardResultsSetPagination
    queryset = AlternativeItem.objects.select_related("item", "alternative_item").all()
    serializer_class = AlternativeItemSerializer
    # permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ["item"]


class RelatedEquipmentViewSet(viewsets.ModelViewSet):
    """ViewSet for RelatedEquipment CRUD operations"""

    queryset = RelatedEquipment.objects.select_related("item").all()
    serializer_class = RelatedEquipmentSerializer
    # permission_classes = [IsAuthenticated]
    pagination_class = StandardResultsSetPagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_fields = ["item"]
    search_fields = ["equipment_name", "equipment_code"]


class PendingTransactionViewSet(viewsets.ModelViewSet):
    """ViewSet for PendingTransaction CRUD operations"""

    queryset = PendingTransaction.objects.select_related("item").all()
    serializer_class = PendingTransactionSerializer
    # permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ["item", "transaction_type", "is_resolved"]
    ordering_fields = ["date", "created_at"]
    ordering = ["-date"]
    pagination_class = StandardResultsSetPagination

    @action(detail=True, methods=["post"], url_path="resolve")
    def resolve(self, request, pk=None):
        """Mark pending transaction as resolved"""
        pending = self.get_object()
        pending.is_resolved = True
        pending.save()
        return Response({"status": "resolved"})


class StockTransactionViewSet(viewsets.ModelViewSet):
    """ViewSet for StockTransaction CRUD operations"""

    queryset = StockTransaction.objects.select_related("item").all()
    serializer_class = StockTransactionSerializer
    # permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ["item"]
    ordering_fields = ["date", "created_at"]
    ordering = ["-date"]
    pagination_class = StandardResultsSetPagination

    @action(detail=False, methods=["get"], url_path="by-item/(?P<item_id>[^/.]+)")
    def by_item(self, request, item_id=None):
        """Get all transactions for specific item"""
        transactions = self.queryset.filter(item_id=item_id).order_by("date")
        serializer = self.get_serializer(transactions, many=True)
        return Response(serializer.data)

class DescriptionViewSet(viewsets.ModelViewSet):
    """ViewSet for Description CRUD operations"""

    queryset = Description.objects.select_related(
        "class_id", "type_id", "gpsc_id"
    ).all()
    serializer_class = DescriptionSerializer
    # permission_classes = [IsAuthenticated]
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    ]
    filterset_fields = ["class_id", "type_id", "gpsc_id"]
    search_fields = ["Des_name", "Des_id", "keyword"]
    ordering_fields = ["created_at", "Des_id"]
    ordering = ["class_id", "type_id", "Des_id"]
    pagination_class = StandardResultsSetPagination

    def get_serializer_class(self):
        # Use the same serializer for both list and detail views
        return DescriptionSerializer
