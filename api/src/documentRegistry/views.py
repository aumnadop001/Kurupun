from rest_framework import viewsets, filters
from rest_framework.permissions import AllowAny
from rest_framework.pagination import PageNumberPagination
from rest_framework.decorators import api_view, permission_classes
from django_filters.rest_framework import DjangoFilterBackend
import django_filters
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
from .models import DocumentRegistry, InventoryRecord
from .serializers import DocumentRegistrySerializer, InventoryRecordSerializer


class DocumentRegistryPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


class DocumentRegistryFilter(django_filters.FilterSet):
    registry_number = django_filters.CharFilter(lookup_expr="icontains")
    sender = django_filters.CharFilter(lookup_expr="icontains")

    class Meta:
        model = DocumentRegistry
        fields = ["registry_number", "sender"]


class DocumentRegistryViewSet(viewsets.ModelViewSet):
    queryset = DocumentRegistry.objects.all()
    serializer_class = DocumentRegistrySerializer
    permission_classes = [AllowAny]
    pagination_class = DocumentRegistryPagination
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    ]
    filterset_class = DocumentRegistryFilter
    search_fields = [
        "registry_number",
        "document_title",
        "sender",
        "related_document_number",
        "withdrawal_set_number",
    ]
    ordering_fields = ["registration_date", "storage_date", "registry_number"]
    ordering = ["-registration_date"]


class InventoryRecordPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


class InventoryRecordFilter(django_filters.FilterSet):
    document_registry = django_filters.NumberFilter(field_name="document_registry__id")

    class Meta:
        model = InventoryRecord
        fields = ["document_registry"]


class InventoryRecordViewSet(viewsets.ModelViewSet):
    queryset = InventoryRecord.objects.all()
    serializer_class = InventoryRecordSerializer
    permission_classes = [AllowAny]
    pagination_class = InventoryRecordPagination
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    ]
    filterset_class = InventoryRecordFilter
    search_fields = [
        "order_criteria",
        "related_equipment",
        "remark",
    ]
    ordering_fields = ["created_at", "doc_date_left", "doc_date_right"]
    ordering = ["-created_at"]


@api_view(["GET"])
@permission_classes([AllowAny])
def export_inventory_excel(request, document_registry_id):
    """Export inventory records to Excel with the format from example.html"""
    try:
        # Fetch document registry
        doc_registry = DocumentRegistry.objects.get(id=document_registry_id)

        # Fetch all inventory records for this document registry
        inventory_records = InventoryRecord.objects.filter(
            document_registry=doc_registry
        ).order_by("created_at")

        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "บันทึกรายการพัสดุ"

        # Define styles
        bold_font = Font(bold=True)
        center_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        header_fill = PatternFill(
            start_color="CCCCCC", end_color="CCCCCC", fill_type="solid"
        )

        # Set column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 14
        ws.column_dimensions["G"].width = 12
        ws.column_dimensions["H"].width = 10
        ws.column_dimensions["I"].width = 12
        ws.column_dimensions["J"].width = 10
        ws.column_dimensions["K"].width = 15
        ws.column_dimensions["L"].width = 20
        ws.column_dimensions["M"].width = 10
        ws.column_dimensions["N"].width = 10
        ws.column_dimensions["O"].width = 10
        ws.column_dimensions["P"].width = 10
        ws.column_dimensions["Q"].width = 10
        ws.column_dimensions["R"].width = 15

        # Header Row 1 (no title row)
        row = 1
        headers_row1 = [
            ("A1", "ว.ด.ป.", 2, 1),  # rowspan=2, colspan=1
            ("B1", "หลักฐาน", 2, 1),
            ("C1", "หมายเลขพัสดุ", 1, 2),
            ("E1", doc_registry.registry_number, 1, 4),
            (
                "I1",
                f"{doc_registry.registry_number} - {doc_registry.document_title}",
                2,
                6,
            ),
            ("O1", "หน่วยนับ", 2, 2),
            ("Q1", "ที่เก็บ", 2, 2),
        ]

        for cell_coord, text, rowspan, colspan in headers_row1:
            cell = ws[cell_coord]
            cell.value = text
            cell.font = bold_font
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = header_fill

        # Header Row 2
        ws["C2"] = "วัน"
        ws["D2"] = "จำนวน"
        ws["E2"] = "หมายเลขพัสดุแทนกันได้"
        ws.merge_cells("E2:H2")

        for col in ["C", "D", "E"]:
            cell = ws[f"{col}2"]
            cell.font = bold_font
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = header_fill

        # Merge cells for rowspan=2
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:B2")
        ws.merge_cells("C1:D1")
        ws.merge_cells("E1:H1")
        ws.merge_cells("I1:N2")
        ws.merge_cells("O1:P2")
        ws.merge_cells("Q1:R2")

        # Info rows (เกณฑ์สั่ง, จุดสั่งเพิ่มเติม, เกณฑ์ปลอดภัย)
        row = 3
        info_labels = ["เกณฑ์สั่ง", "จุดสั่งเพิ่มเติม", "เกณฑ์ปลอดภัย"]
        for label in info_labels:
            ws.merge_cells(f"A{row}:B{row}")
            cell = ws[f"A{row}"]
            cell.value = label
            cell.font = bold_font
            cell.alignment = left_alignment
            cell.border = thin_border
            row += 1

        # Merge cells for info section
        ws.merge_cells("E3:H5")
        ws["E3"].value = "หมายเลขพัสดุแทนกันได้"
        ws["E3"].alignment = center_alignment
        ws["E3"].border = thin_border

        ws.merge_cells("I3:N5")
        ws["I3"].value = "ครุภัณฑ์ที่เกี่ยวข้อง"
        ws["I3"].alignment = center_alignment
        ws["I3"].border = thin_border

        ws.merge_cells("O3:R5")
        ws["O3"].value = "หมายเหตุ"
        ws["O3"].alignment = center_alignment
        ws["O3"].border = thin_border

        # Main headers
        row = 6
        ws.merge_cells(f"A{row}:H{row}")
        ws[f"A{row}"] = "ค้างรับ และ ค้างจ่าย"
        ws[f"A{row}"].font = bold_font
        ws[f"A{row}"].alignment = center_alignment
        ws[f"A{row}"].border = thin_border
        ws[f"A{row}"].fill = header_fill

        ws.merge_cells(f"I{row}:R{row}")
        ws[f"I{row}"] = "ความต้องการรับและจ่าย"
        ws[f"I{row}"].font = bold_font
        ws[f"I{row}"].alignment = center_alignment
        ws[f"I{row}"].border = thin_border
        ws[f"I{row}"].fill = header_fill

        # Column headers row 7-8
        row = 7
        column_headers = [
            ("A", "ว.ด.ป.", 2),
            ("B", "หลักฐาน", 2),
            ("C", "หน่วย", 2),
            ("D", "จำนวน", 2),
            ("E", "รับ\nค้าง", 2),
            ("F", "รับ\nค้าง", 2),
            ("G", "รับ\nค้าง", 2),
            ("H", "รับ\nค้าง", 2),
            ("I", "ว.ด.ป.", 2),
            ("J", "รับ", 2),
            ("K", "ราคาต่อหน่วย", 2),
            ("L", "หลักฐาน", 2),
            ("M", "ความต้องการ\nขั้นต้น", 1),
            ("N", "ความต้องการ\nทดแทน", 1),
            ("O", "จ่าย", 2),
            ("P", "รวมยืม", 2),
            ("Q", "คงคลัง", 2),
            ("R", "ลายมือชื่อ", 2),
        ]

        for col, text, rowspan in column_headers:
            cell = ws[f"{col}{row}"]
            cell.value = text
            cell.font = bold_font
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = header_fill
            if rowspan == 2:
                ws.merge_cells(f"{col}{row}:{col}{row+1}")

        # Merge for "ความต้องการ" header
        ws.merge_cells(f"M{row}:N{row}")
        ws[f"M{row}"] = "ความต้องการ"
        ws[f"M{row}"].font = bold_font
        ws[f"M{row}"].alignment = center_alignment
        ws[f"M{row}"].border = thin_border
        ws[f"M{row}"].fill = header_fill

        # Sub-headers for row 8
        ws["M8"] = "ขั้นต้น"
        ws["M8"].font = bold_font
        ws["M8"].alignment = center_alignment
        ws["M8"].border = thin_border
        ws["M8"].fill = header_fill

        ws["N8"] = "ทดแทน"
        ws["N8"].font = bold_font
        ws["N8"].alignment = center_alignment
        ws["N8"].border = thin_border
        ws["N8"].fill = header_fill

        # Data rows
        row = 9
        for record in inventory_records:
            # Left section (ค้างรับ และ ค้างจ่าย)
            ws[f"A{row}"] = (
                record.doc_date_left.strftime("%d/%m/%Y")
                if record.doc_date_left
                else ""
            )
            ws[f"B{row}"] = record.evidence_left or ""
            ws[f"C{row}"] = record.unit_left or ""
            ws[f"D{row}"] = record.qty_left if record.qty_left else ""
            ws[f"E{row}"] = record.pending_receive_1 if record.pending_receive_1 else ""
            ws[f"F{row}"] = record.pending_receive_2 if record.pending_receive_2 else ""
            ws[f"G{row}"] = record.pending_receive_3 if record.pending_receive_3 else ""
            ws[f"H{row}"] = record.pending_receive_4 if record.pending_receive_4 else ""

            # Right section (ความต้องการรับและจ่าย)
            ws[f"I{row}"] = (
                record.doc_date_right.strftime("%d/%m/%Y")
                if record.doc_date_right
                else ""
            )
            ws[f"J{row}"] = record.received_qty if record.received_qty else ""
            ws[f"K{row}"] = (
                float(record.price_per_unit) if record.price_per_unit else ""
            )
            ws[f"L{row}"] = record.evidence_right or ""
            ws[f"M{row}"] = record.demand_initial if record.demand_initial else ""
            ws[f"N{row}"] = record.demand_replace if record.demand_replace else ""
            ws[f"O{row}"] = record.issued_qty if record.issued_qty else ""
            ws[f"P{row}"] = record.total_borrowed if record.total_borrowed else ""
            ws[f"Q{row}"] = record.stock_balance if record.stock_balance else ""
            ws[f"R{row}"] = record.signature or ""

            # Apply borders and alignment
            for col in [
                "A",
                "B",
                "C",
                "D",
                "E",
                "F",
                "G",
                "H",
                "I",
                "J",
                "K",
                "L",
                "M",
                "N",
                "O",
                "P",
                "Q",
                "R",
            ]:
                cell = ws[f"{col}{row}"]
                cell.border = thin_border
                if col in ["D", "E", "F", "G", "H", "J", "K", "M", "N", "O", "P", "Q"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            row += 1

        # Create HTTP response
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        current_date = datetime.now().strftime("%Y%m%d")
        filename = f"บันทึกรายการพัสดุ_{doc_registry.registry_number}_{current_date}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        wb.save(response)
        return response

    except DocumentRegistry.DoesNotExist:
        return HttpResponse("Document Registry not found", status=404)
    except Exception as e:
        return HttpResponse(f"Error: {str(e)}", status=500)


def export_excel(request):
    # สร้าง Workbook และเลือก Sheet แรก
    wb = Workbook()
    ws = wb.active
    ws.title = "Documents"

    # เขียนหัวตาราง
    ws.append(["ทะเบียนที่", "วันที่ลงทะเบียน", "เอกสาร", "จาก", "รายการแรก"])

    # ดึงข้อมูลจากฐานข้อมูล
    for doc in Document.objects.all():
        ws.append(
            [
                doc.registry_number,
                doc.registration_date.strftime("%d/%m/%Y"),
                doc.document_title,
                doc.sender,
                doc.first_item,
            ]
        )

    # สร้าง Response เพื่อให้ browser ดาวน์โหลดไฟล์
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    currentDate = datetime.now().strftime("%Y-%m-%d")
    filename = f"ทะเบียนเอกสาร_{currentDate}.xlsx"
    # filename = "{}_{}.xlsx".format("ทะเบียนคุมเอกสาร", currentDate)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response
